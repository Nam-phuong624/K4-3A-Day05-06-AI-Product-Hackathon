import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

excel_path = "eval/Bảng so sánh 20 test case trước - sau.xlsx"
wb = openpyxl.load_workbook(excel_path)

# Styles
font_warning = Font(name="Calibri", size=9, bold=True, color="9C0006")
fill_warning = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
font_success = Font(name="Calibri", size=9, bold=True, color="006100")
fill_success = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
font_data = Font(name="Calibri", size=9, color="000000")
font_bold = Font(name="Calibri", size=9, bold=True, color="000000")

thin_border = Border(
    left=Side(style='thin', color='BFBFBF'),
    right=Side(style='thin', color='BFBFBF'),
    top=Side(style='thin', color='BFBFBF'),
    bottom=Side(style='thin', color='BFBFBF')
)

align_top_left = Alignment(horizontal="left", vertical="top", wrap_text=True)
align_top_center = Alignment(horizontal="center", vertical="top", wrap_text=True)

# 1. TẠO MỚI HOÀN TOÀN SHEET 3: Case_Study_Live_K4_PAIR
sheet_title = "Case_Study_Live_K4_PAIR"
if sheet_title in wb.sheetnames:
    del wb[sheet_title]
ws3 = wb.create_sheet(title=sheet_title)
ws3.views.sheetView[0].showGridLines = True

# Title Block
ws3.merge_cells("A1:K1")
ws3["A1"] = "CASE STUDY THỰC NGHIỆM: ĐỐI ĐẦU CHUỖI ĐÀO SÂU SOCRATIC VỀ GOOGLE PAIR TRÊN BÀI HỌC DAY 02"
ws3["A1"].font = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
ws3["A1"].fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
ws3["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws3.row_dimensions[1].height = 36

ws3.merge_cells("A2:K2")
ws3["A2"] = "Phân tích Lỗ hổng Trích dẫn: Nguồn ở box chân tin nhắn là THỰC, các số trang tự phát sinh trong câu văn không có ở box là NGUỒN BỊA (Phantom Citation)"
ws3["A2"].font = Font(name="Calibri", size=10, italic=True, color="595959")
ws3["A2"].alignment = Alignment(horizontal="center", vertical="center")
ws3.row_dimensions[2].height = 20

columns_s3_new = [
    ("Tầng Đào Sâu (Turn)", 12, align_top_center),
    ("Câu Hỏi Đào Sâu", 24, align_top_left),
    ("Phản Hồi Chatbot Thật Của Lớp (VLearn Production)", 35, align_top_left),
    ("Nguồn Viết Trong Text (Inline Citations)", 20, align_top_left),
    ("Nguồn Thực Ở Box Chân Tin Nhắn (Footer Citation)", 22, align_top_left),
    ("Bản Chất Lỗi Bịa Nguồn (Phantom Citation Audit)", 26, align_top_left),
    ("Giải Pháp AI Tutor Mới Đề Xuất (Clean & Grounded)", 32, align_top_left),
    ("Độ Dài Mới", 9, align_top_center),
    ("Nguồn Xác Minh Độc Lập Chuẩn Xác", 22, align_top_left),
    ("Gợi Mở Socratic (Option A/B)", 25, align_top_left),
    ("Giá Trị UX Sư Phạm Vượt Trội", 26, align_top_left)
]

# Header Grouping Row 4
ws3.merge_cells("A4:F4")
ws3.cell(row=4, column=1, value="HIỆN TRẠNG THỰC TẾ TRÊN CHATBOT THẬT CỦA LỚP K4: ĐỐI CHIẾU LỖI BỊA NGUỒN TRONG TEXT VS NGUỒN THẬT").font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
ws3.cell(row=4, column=1).fill = PatternFill(start_color="843C0C", end_color="843C0C", fill_type="solid")
ws3.cell(row=4, column=1).alignment = Alignment(horizontal="center", vertical="center")
for c in range(1, 7):
    ws3.cell(row=4, column=c).border = thin_border
    ws3.cell(row=4, column=c).fill = PatternFill(start_color="843C0C", end_color="843C0C", fill_type="solid")

ws3.merge_cells("G4:K4")
ws3.cell(row=4, column=7, value="GIẢI PHÁP ĐỀ XUẤT MỚI: TRIỆT TIÊU 100% BỊA NGUỒN & SẠCH TEXT").font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
ws3.cell(row=4, column=7).fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
ws3.cell(row=4, column=7).alignment = Alignment(horizontal="center", vertical="center")
for c in range(7, 12):
    ws3.cell(row=4, column=c).border = thin_border
    ws3.cell(row=4, column=c).fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
ws3.row_dimensions[4].height = 24

# Headers Row 5
for col_idx, (c_name, c_width, _) in enumerate(columns_s3_new, 1):
    c = ws3.cell(row=5, column=col_idx, value=c_name)
    c.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    c.fill = PatternFill(start_color="262626", end_color="262626", fill_type="solid")
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = thin_border
    ws3.column_dimensions[get_column_letter(col_idx)].width = c_width
ws3.row_dimensions[5].height = 32

pair_audit_data = [
    {
        "turn": "Tầng 1 (Root)",
        "q": "pair là gì",
        "reply": "Chào Nam, trong nội dung bài học hôm nay, PAIR là viết tắt của Google People + AI Guidebook trang 6...\nChương 1: trang 6 trang 125.\nChương 6: trang 6 trang 146 trang 147.\nPM trả lời câu hỏi: trang 3 trang 6.",
        "inline_cite": "trang 6, trang 125, trang 146, trang 147, trang 3",
        "footer_cite": "k4-d2-slide-blue-Slide D301:\n• Tr. 6 (3 bước quyết định)\n• Tr. 3 (3 câu hỏi ở trên)",
        "audit_flaw": "BỊA NGUỒN NGHIÊM TRỌNG (Phantom Citation):\n• Nguồn box chân chỉ có Tr. 3 và Tr. 6.\n• Trong text bot tự bịa 'trang 125', 'trang 146', 'trang 147' hoàn toàn không có ở box nguồn (Slide chỉ có ~33 trang)!",
        "new_reply": "PAIR (Google People + AI Guidebook) là bộ cẩm nang thiết kế sản phẩm AI lấy con người làm trung tâm, tập trung vào 2 trọng tâm: Xác định nhu cầu người dùng và Xử lý lỗi hệ thống an toàn, minh bạch.",
        "new_len": 198,
        "new_cite": "Slide [k4-d2-slide-blue] · Trang 6",
        "socratic": "• Option A: 3 bước quyết định có nên dùng AI theo PAIR là gì?\n• Option B: Cách PAIR hướng dẫn thiết kế trải nghiệm khi AI xảy ra lỗi?",
        "benefit": "Khớp 100% với box nguồn chuẩn Trang 6; loại bỏ hoàn toàn số trang ma bịa đặt (125, 146, 147)."
    },
    {
        "turn": "Tầng 2 (Quy trình)",
        "q": "3 bước quyết định AI theo PAIR gồm những gì?",
        "reply": "Chào Nam, theo Google PAIR, 3 bước quyết định có nên dùng AI hay không bao gồm trang 28:\n- Giao điểm: trang 28.\n- Automate vs Augment: trang 28.\n- Reward function: trang 28.\n- Trả lời 3 câu hỏi: trang 28.",
        "inline_cite": "Lặp lại 'trang 28' 5 lần trong văn bản",
        "footer_cite": "k4-d2-slide-blue-Slide D301:\n• Tr. 28 (Bài toán có nằm trong nhóm AI làm tốt hơn rule không?)",
        "audit_flaw": "Ô NHIỄM VĂN BẢN (Text Pollution):\n• Nguồn Tr. 28 ở box dưới là chính xác, nhưng bot nhồi nhét 'trang 28' vào từng câu trong text gây vụn vặt và mất mạch đọc của học viên.",
        "new_reply": "Theo PAIR, 3 bước quyết định gồm: (1) Xác định giao điểm nhu cầu người dùng và thế mạnh AI; (2) Lựa chọn mức độ tự động hóa (Automate vs Augment); và (3) Thiết lập Reward Function cùng tiêu chí thành công đo lường được.",
        "new_len": 218,
        "new_cite": "Slide [k4-d2-slide-blue] · Trang 28",
        "socratic": "• Option A: Lợi thế của AI khác gì so với phương pháp Rule/Heuristic truyền thống?\n• Option B: Khi nào nên chọn Augment thay vì Automate hoàn toàn?",
        "benefit": "Giữ đúng nguồn chuẩn Trang 28 ở Badge riêng; làm sạch 100% từ rác 'trang 28' trong câu trả lời."
    },
    {
        "turn": "Tầng 3 (Bản chất)",
        "q": "Lợi thế của AI khác gì so với Rule/Heuristic?",
        "reply": "Chào Nam, theo tài liệu hướng dẫn của PAIR, sự khác biệt giữa AI và Rule/Heuristic được thể hiện qua 3 khía cạnh... trang 29... trang 33... trang 29.",
        "inline_cite": "trang 29, trang 33, trang 29",
        "footer_cite": "k4-d2-slide-blue-Slide D301:\n• Tr. 29 (Bảng so sánh Rule vs AI)\n• Tr. 33 (Lựa chọn tối ưu)",
        "audit_flaw": "TRÍCH DẪN KHỚP NGUỒN NHƯNG BỊ BỨC TƯỜNG CHỮ:\n• Nguồn Tr. 29 và Tr. 33 là chính xác, nhưng phản hồi quá dài (1.142 ký tự), dính lỗi chèn 'trang 29/33' vào text.",
        "new_reply": "Khác biệt cốt lõi: Rule tối ưu cho đầu vào đếm được, có 1 đáp án duy nhất và chi phí sai rất đắt; trong khi AI vượt trội khi đầu vào vô hạn biến thiên ngữ cảnh và hệ thống có khả năng học hỏi, dung nạp sai số.",
        "new_len": 210,
        "new_cite": "Slide [k4-d2-slide-blue] · Trang 29, Trang 33",
        "socratic": "• Option A: Tại sao cần xác định đối tượng (Actor) trước khi tìm bài toán AI?\n• Option B: Làm thế nào để xác định đầu vào của bài toán theo dữ liệu thực tế?",
        "benefit": "Tóm gọn bản chất 2 dòng; hiển thị trích dẫn đúng Trang 29 & 33 ở chân thẻ."
    },
    {
        "turn": "Tầng 4 (Vòng lặp)",
        "q": "Làm thế nào để xác định đầu vào của bài toán?",
        "reply": "Chào Nam, để xác định đầu vào... dựa vào ba câu hỏi chẩn đoán trong bảng so sánh trang 29... Đếm được hay vô hạn trang 29... Có đáp án duy nhất trang 29... Mức độ rủi ro trang 29... Kỹ thuật phân kỳ trang 18.",
        "inline_cite": "trang 29 (lặp 4 lần), trang 18",
        "footer_cite": "k4-d2-slide-blue-Slide D301:\n• Tr. 29 (Bảng so sánh Rule vs AI)\n• Tr. 18 (Phân kỳ để thấu hiểu)",
        "audit_flaw": "VÒNG LẶP NGỮ NGHĨA & NGUỒN TRÙNG LẶP:\n• Nguồn box chân đúng Tr. 18 & Tr. 29, nhưng trong text bot nhại lại nguyên vẹn Tr. 29 của Tầng 3, không mở rộng được kiến thức mới về phương pháp phân tích đầu vào.",
        "new_reply": "Để xác định đầu vào, áp dụng kỹ thuật phân kỳ: Quan sát thực tế (Field Study), nhật ký hành vi (Diary Study) và rà soát log hệ thống để phân loại đầu vào thuộc dạng cấu trúc cố định (dùng Rule) hay ngôn ngữ tự nhiên mở (dùng AI).",
        "new_len": 230,
        "new_cite": "Slide [k4-d2-slide-blue] · Trang 18",
        "socratic": "• Option A: Tại sao cần xác định đối tượng (Actor) trước khi phân tích đầu vào quy trình?\n• Option B: Kỹ thuật Field Study và Diary Study khác nhau như thế nào khi khảo sát người dùng?",
        "benefit": "Tập trung chuẩn xác vào Trang 18 (Field/Diary Study), chấm dứt việc nhại lại nguồn Trang 29."
    },
    {
        "turn": "Tầng 5 (Đối tượng)",
        "q": "Tại sao cần xác định đối tượng trước khi tìm bài toán?",
        "reply": "Chào Nam, việc xác định đối tượng là bước tiên quyết... Quyết định Workflow trang 8... Xác định Metric trang 8... Đánh giá rủi ro trang 8 trang 109... Xây dựng tính năng trang 8.",
        "inline_cite": "trang 8 (lặp 4 lần), trang 109",
        "footer_cite": "k4-d2-slide-blue-Slide D301:\n• Tr. 8 (Đối tượng khác nhau dẫn đến workflow, metric, rủi ro khác nhau)",
        "audit_flaw": "BỊA NGUỒN TRANG 109 (Phantom Citation):\n• Box nguồn ở dưới chỉ có DUY NHẤT 'Tr. 8'.\n• Trong text bot tự bịa thêm 'trang 109' hoàn toàn không có trong tài liệu nguồn bài giảng (Học viên click vào bị lỗi 404)!",
        "new_reply": "Xác định đối tượng là tiền đề bắt buộc vì: Đối tượng khác nhau sẽ quyết định Workflow khác nhau, Metric thành công khác nhau (CSAT vs Thời gian xử lý) và Nút thắt rủi ro khác nhau, tránh việc làm tính năng công nghệ mà không ai dùng.",
        "new_len": 232,
        "new_cite": "Slide [k4-d2-slide-blue] · Trang 8",
        "socratic": "• Option A: Làm sao để xác định nút thắt trong quy trình của đối tượng?\n• Option B: Các chỉ số (Metrics) đo lường thành công khác biệt thế nào giữa user nội bộ và khách hàng?",
        "benefit": "Xóa bỏ triệt để nguồn bịa 'trang 109'; neo duy nhất và chuẩn xác vào Slide Trang 8."
    },
    {
        "turn": "Tầng 6 (Nút thắt)",
        "q": "Làm sao để xác định nút thắt trong quy trình của đối tượng?",
        "reply": "Để xác định nút thắt... trang 18... trang 19... trang 12... trang 24... trang 18... trang 12 trang 24...",
        "inline_cite": "trang 18, trang 19, trang 12, trang 24 (rải rác khắp bài)",
        "footer_cite": "k4-d2-slide-blue-Slide D301:\n• Tr. 18 (Field Study)\n• Tr. 19 (Việc lặp lại, ngồi chờ)\n• Tr. 12 (4 câu hỏi chẩn đoán)\n• Tr. 24 (Báo cáo tuần: viết narrative 25')",
        "audit_flaw": "BỨC TƯỜNG CHỮ CỰC ĐẠI & KHÔNG BẤM ĐƯỢC:\n• Box nguồn ở dưới trích dẫn rất chuẩn 4 trang (12, 18, 19, 24).\n• Tuy nhiên trong text bot xả tới 1.820 ký tự và biến 2 câu hỏi gợi mở thành text chết, không hỗ trợ bấm nút tương tác.",
        "new_reply": "Xác định nút thắt bằng 3 bước: (1) Chia nhỏ quy trình theo mốc thời gian (ví dụ báo cáo tuần: viết narrative chiếm 25' là nút thắt); (2) Đặt 4 câu hỏi chẩn đoán để tìm bước quá tải; và (3) Áp dụng kỹ thuật 5 Whys truy nguồn gốc rễ.",
        "new_len": 235,
        "new_cite": "Slide [k4-d2-slide-blue] · Trang 12, Trang 18, Trang 24",
        "socratic": "• Option A: Cách dùng ma trận Tác động – Nỗ lực (Impact-Effort Matrix) để chọn bài toán?\n• Option B: Làm sao để lượng hóa điểm đau (Pain Point) của người dùng thành con số?",
        "benefit": "Thu gọn từ 1.820 xuống 235 ký tự; giữ nguyên trích dẫn 3 nguồn cốt lõi (12, 18, 24); biến 2 câu text chết thành nút bấm tương tác."
    }
]

r = 6
for item in pair_audit_data:
    row_vals = [
        item["turn"],
        item["q"],
        item["reply"],
        item["inline_cite"],
        item["footer_cite"],
        item["audit_flaw"],
        item["new_reply"],
        item["new_len"],
        item["new_cite"],
        item["socratic"],
        item["benefit"]
    ]
    for col_i, val in enumerate(row_vals, 1):
        cell = ws3.cell(row=r, column=col_i, value=val)
        cell.font = font_data
        cell.border = thin_border
        cell.alignment = columns_s3_new[col_i - 1][2]
        
        if col_i == 1:
            cell.font = font_bold
        elif col_i == 6:
            if "BỊA NGUỒN" in str(val):
                cell.fill = fill_warning
                cell.font = font_warning
        elif col_i == 8:
            cell.font = font_bold
        elif col_i == 9:
            cell.font = font_success
    ws3.row_dimensions[r].height = 110
    r += 1

wb.save(excel_path)
print("Updated excel successfully with clean MergedCells handling!")
